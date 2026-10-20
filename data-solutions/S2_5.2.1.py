from __future__ import annotations

"""
投资策略执行情况专题 —— V5.2.1 阶段二
PPT文件提取与比例核查

输入：
1. 产品资料\产品清单.xlsx
2. 日报产品比例汇总.xlsx（阶段一输出）
3. 投决会目录下所有层级的.pptx
4. 宣推材料目录下所有层级的.pptx

输出：
- 文件核查表.xlsx
  - 核查结果

核心规则：
- 仅核查产品清单中的产品；
- 投决会目录下的PPT直接认定为投决会文件；
- 宣推材料目录下的PPT直接认定为宣推文件；
- 第一遍使用“产品系列核心词＋单一编号/编号区间”匹配带号产品；
- 第一遍未匹配的文件进入第二遍，对产品清单中本身不含“数字＋号”的产品按清理后完整名称严格匹配；
- 禁止仅凭“2号FOF”等宽泛尾部关键词匹配；
- 文件名已识别为其他产品系列时，禁止正文回退；
- 支持“1-X号”区间文件在同一产品系列内一对多匹配；
- 文件名中“万家共赢”之前的事项编号、附件编号等前缀不参与产品系列识别；
- 文件名中“资产管理计划”之后的投决会说明、日期、版本号和流转备注不参与产品系列识别；
- “均衡3号推介材料”等不完整文件名进入PPT正文严格匹配；
- 宣推历史模板中的“积极型/均衡型策略组合”可作为配置候选页；
- 根据产品名称选择稳健/均衡/积极/指增/增强配置方案；
- 多个同类配置方案全部抄写，不默认选择第一行或第一列；
- 文件策略只写简短策略名称，不写整段介绍；
- 比例必须带%，且必须位于0%-100%；
- 目标客户、产品定位仅从产品要素/风险收益特征候选页提取；
- 不读取或核对任何文件日期；
- V5.1.0改为逐文件打开、处理、释放PPT，不再缓存Presentation和整份页面文本；
- V5.1.2删除目标收益和最大回撤字段，改为提取“目标客户”表头右侧同一行的全部文字。
- V5.1.3将“风险收益特征”加入目标客户字段的表头关键词，二者均抄写同一行右侧全部文字。
- V5.2.0当一级分类无法对应阶段一结果时，支持按文件策略中的“可转债/宏观”回退到科目二级分类汇总核查；若两者同时出现，则合计该产品全部可转债与宏观科目占比。
- V5.2.0原生表格未提取到策略配置时，补充扫描配置页面文本框，识别“策略名称 + x%-y%”，并支持名称与比例分处相邻文本框。
- V5.2.1宣推文件不再预先筛选“资产配置方案/投资方案”页面；与投决会一致，先遍历全部页面的原生表格，表格未提取到策略后再启用文本框补充识别。
- 单个PPT读取或处理异常不会终止整个批次，主程序异常打印完整traceback。
"""

import gc
import re
import sys
import traceback
import unicodedata
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator, Optional, Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.shapes.group import GroupShape

try:
    from colorama import Fore, init as colorama_init
except ImportError:
    class _Ansi:
        GREEN = "\033[32m"
        YELLOW = "\033[33m"
        RED = "\033[31m"
        RESET = "\033[0m"

    Fore = _Ansi()  # type: ignore[assignment]

    def colorama_init(*, autoreset: bool = True) -> None:
        del autoreset


# ============================================================
# 1. 配置
# ============================================================


@dataclass(frozen=True)
class Config:
    project_dir: Path = Path(
        r"C:\Users\gyhg-intern\Desktop\投资策略执行情况专题"
    )

    product_list_relative: Path = Path("产品资料") / "产品清单.xlsx"
    stage1_output_name: str = "日报产品比例汇总.xlsx"
    committee_dir_name: str = "投决会"
    promotion_dir_name: str = "宣推材料"
    output_name: str = "文件核查表.xlsx"

    product_list_header: str = "产品清单"
    stage1_sheet: str = "产品策略分类"
    stage1_detail_sheet: str = "科目比例明细"
    output_sheet: str = "核查结果"

    enable_content_matching: bool = True
    content_match_slide_limit: int = 0  # 0 = 全部页面

    @property
    def product_list_path(self) -> Path:
        return self.project_dir / self.product_list_relative

    @property
    def stage1_output_path(self) -> Path:
        return self.project_dir / self.stage1_output_name

    @property
    def committee_dir(self) -> Path:
        return self.project_dir / self.committee_dir_name

    @property
    def promotion_dir(self) -> Path:
        return self.project_dir / self.promotion_dir_name

    @property
    def output_path(self) -> Path:
        return self.project_dir / self.output_name


CONFIG = Config()

# 文件名和产品名中不用于区分产品的通用字段。
GENERIC_PRODUCT_PARTS = (
    "万家共赢资产管理有限公司",
    "万家共赢",
    "集合资产管理计划",
    "集合资产计划",
    "单一资产管理计划",
    "专项资产管理计划",
    "集合资管计划",
    "单一资管计划",
    "资产管理计划",
    "宣传推介材料附件",
    "宣传推介材料",
    "宣传推介",
    "宣推材料",
    "推介材料",
    "投决会上会材料",
    "投决会材料",
    "上会材料",
    "投资决策委员会",
    "投资决策会",
    "投决会",
    "汇报材料",
    "产品材料",
    "最终版",
    "修订版",
    "定稿版",
    "上会版",
    "最新版",
)


CATEGORY_ORDER = (
    "量化选股",
    "指数增强",
    "股票多头",
    "相对价值",
    "市场中性",
    "固定收益",
    "宏观策略",
    "管理期货",
)

# 更具体的别名必须放在前面。
STRATEGY_ALIASES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("股票多头", ("股票多头", "量化多头", "多头策略", "多头")),
    ("指数增强", ("中证全指指数增强", "量化指增", "指数增强", "指增")),
    ("量化选股", ("量化选股", "股票量化", "量化策略")),
    ("相对价值", ("相对价值", "套利策略", "套利")),
    ("市场中性", ("市场中性", "股票中性", "中性策略", "中性")),
    ("固定收益", ("固定收益", "固收＋", "固收+", "固收")),
    ("宏观策略", ("宏观策略", "宏观")),
    ("管理期货", ("管理期货", "量化CTA", "CTA")),
)

# 产品名称中的方案关键词与配置表中的方案别名。
SCHEME_ALIASES: dict[str, tuple[str, ...]] = {
    "稳健": ("稳健", "稳健型", "稳健系列"),
    "均衡": ("均衡", "均衡型", "均衡系列"),
    "积极": ("积极", "积极型", "积极系列"),
    "指增": ("指增", "指增型", "指增系列", "指数增强型", "指数增强系列"),
    "增强": ("增强", "增强型", "增强系列"),
}

# 避免“增强”抢先命中“指增/指数增强”。
SCHEME_DETECTION_ORDER = ("指增", "稳健", "均衡", "积极", "增强")

ALLOCATION_HEADER_KEYWORDS = (
    "资产配置比例",
    "配置参考比例",
    "策略配置比例",
    "策略配置参考比例",
)

PROMOTION_PAGE_KEYWORDS = ("资产配置方案", "投资方案")
PROMOTION_SCHEME_PAGE_KEYWORDS = ("策略组合", "策略配置组合")
METADATA_PAGE_KEYWORDS = ("产品要素", "风险收益特征")
METADATA_AUXILIARY_PAGE_KEYWORDS = ("目标客户",)
TARGET_CUSTOMER_HEADERS = ("目标客户", "风险收益特征")

METADATA_FIELDS: dict[str, tuple[str, ...]] = {
    "product_positioning": ("产品定位", "产品类型"),
}

GENERIC_TABLE_LABELS = {
    "集合产品系列",
    "定制产品系列",
    "产品系列",
    "产品类型",
    "策略类型",
    "策略种类",
    "策略类别",
    "大类策略",
    "资产配置比例",
    "配置参考比例",
    "策略配置比例",
    "策略配置参考比例",
    "比例",
    "占比",
}

HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
RED_FILL = PatternFill("solid", fgColor="F4CCCC")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")

INVISIBLE_CHARS_RE = re.compile(r"[\u200b-\u200f\u202a-\u202e\ufeff]")
SEPARATOR_RE = re.compile(r"[\s_\-—–·•,，。；;：:（）()\[\]【】<>《》/\\]+")
STRATEGY_RATIO_PATTERN = re.compile(r"策略.{0,12}(?:比例|占比)", re.IGNORECASE)

SECONDARY_FALLBACK_KEYWORDS = ("可转债", "宏观")

# 文本框补充识别只接受完整的 x%-y% 区间。
TEXT_RATIO_PATTERN = re.compile(
    r"(?P<lower>\d{1,3}(?:\.\d+)?)\s*[%％]\s*"
    r"(?:-|—|–|－|~|～|至)\s*"
    r"(?P<upper>\d{1,3}(?:\.\d+)?)\s*[%％]",
    re.IGNORECASE,
)

TEXT_ALLOCATION_PAGE_KEYWORDS = (
    "投资方案",
    "资产配置方案",
    "策略组合",
    "策略配置组合",
    "策略配置",
    "资产配置",
    "配置比例",
)

TEXT_NON_STRATEGY_LABELS = {
    *GENERIC_TABLE_LABELS,
    *TEXT_ALLOCATION_PAGE_KEYWORDS,
    *ALLOCATION_HEADER_KEYWORDS,
    "产品配置",
    "配置方案",
    "组合配置",
    "策略配置结构",
    "本产品近期拟定的策略配置结构如下图所示",
}


# ============================================================
# 2. 数据结构
# ============================================================


@dataclass(frozen=True)
class ProductRecord:
    name: str
    order: int


@dataclass(frozen=True)
class ActualStrategyRecord:
    product_name: str
    primary_category: str
    total_ratio: float


@dataclass(frozen=True)
class SecondaryStrategyRecord:
    product_name: str
    secondary_category: str
    subject_ratio: float


@dataclass(frozen=True)
class ProductIdentity:
    """产品清单中单一母产品的结构化身份。"""

    product_name: str
    prefix: str
    suffix: str
    number: int
    order: int

    @property
    def series_core(self) -> str:
        return f"{self.prefix}{self.suffix}"

    @property
    def full_token(self) -> str:
        return f"{self.prefix}{self.number}号{self.suffix}"


@dataclass(frozen=True)
class NumberlessProductIdentity:
    """产品清单中不含“数字＋号”的母产品身份。"""

    product_name: str
    full_token: str
    order: int


@dataclass(frozen=True)
class FileIdentity:
    """从PPT文件名中解析出的产品系列及单一编号/编号区间。"""

    prefix: str
    suffix: str
    start_number: int
    end_number: int
    is_range: bool

    @property
    def series_core(self) -> str:
        return f"{self.prefix}{self.suffix}"

    @property
    def range_label(self) -> str:
        if self.is_range:
            return f"{self.start_number}-{self.end_number}号"
        return f"{self.start_number}号"

    def contains(self, number: int) -> bool:
        low, high = sorted((self.start_number, self.end_number))
        return low <= number <= high


@dataclass(frozen=True)
class PageMatchResult:
    product_names: tuple[str, ...] = ()
    is_range: bool = False
    ambiguous: bool = False
    detail: str = ""


@dataclass(frozen=True)
class PptFile:
    path: Path
    file_type: str


@dataclass(frozen=True)
class TextShapeRecord:
    text: str
    left: int
    top: int
    width: int
    height: int
    order: int

    @property
    def right(self) -> int:
        return self.left + self.width

    @property
    def bottom(self) -> int:
        return self.top + self.height

    @property
    def center_x(self) -> float:
        return self.left + self.width / 2

    @property
    def center_y(self) -> float:
        return self.top + self.height / 2


@dataclass(frozen=True)
class ExtractedPlan:
    raw_strategy: str
    primary_category: str
    lower_limit: Optional[float]
    upper_limit: Optional[float]
    slide_number: int
    scheme_label: str
    extraction_status: str
    source_order: int


@dataclass(frozen=True)
class PptMetadata:
    target_customer: str = ""
    product_positioning: str = ""


@dataclass(frozen=True)
class ExtractionResult:
    plans: tuple[ExtractedPlan, ...]
    status: str


@dataclass
class CheckRow:
    product_name: str
    primary_category: str
    file_name: str
    file_type: str
    file_strategy: str
    actual_total_ratio: Optional[float]
    lower_limit: Optional[float]
    upper_limit: Optional[float]
    check_result: str
    target_customer: str = ""
    product_positioning: str = ""

    # 仅用于排序，不写入Excel。
    product_order: int = 0
    source_order: int = 0


@dataclass
class RunStats:
    product_count: int = 0
    actual_strategy_count: int = 0
    secondary_strategy_count: int = 0
    committee_file_count: int = 0
    promotion_file_count: int = 0
    matched_files: set[Path] = field(default_factory=set)
    unmatched_files: int = 0
    products_without_committee: int = 0
    products_without_promotion: int = 0
    extracted_plans: int = 0
    completed_checks: int = 0
    compliant: int = 0
    below: int = 0
    above: int = 0
    unmatched_actual_strategy: int = 0
    manual_review: int = 0
    abnormal_ratio: int = 0
    secondary_fallback_checks: int = 0
    text_fallback_plans: int = 0


# ============================================================
# 3. 日志
# ============================================================


class Logger:
    def __init__(self) -> None:
        colorama_init(autoreset=True)

    @staticmethod
    def _line(
        stage: str,
        product: str = "",
        strategy: str = "",
        file_name: str = "",
        status: str = "",
    ) -> str:
        return (
            f"[{stage}] 产品：{product or 'N/A'} | "
            f"科目/策略：{strategy or 'N/A'} | "
            f"文件：{file_name or 'N/A'} | 状态：{status}"
        )

    def info(self, *args: str) -> None:
        print(self._line(*args))

    def success(self, *args: str) -> None:
        print(Fore.GREEN + self._line(*args))

    def warning(self, *args: str) -> None:
        print(Fore.YELLOW + self._line(*args))

    def error(self, *args: str) -> None:
        print(Fore.RED + self._line(*args))


LOGGER = Logger()


# ============================================================
# 4. 通用清洗
# ============================================================


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    text = INVISIBLE_CHARS_RE.sub("", text)
    text = text.replace("\x0b", "\n")
    text = text.replace("\r", "\n")
    text = re.sub(r"[\t ]+", " ", text)
    text = re.sub(r"\n+", "\n", text)
    return text.strip()


def compact_text(value: Any) -> str:
    return SEPARATOR_RE.sub("", clean_text(value)).lower()


def normalize_range_text(value: Any) -> str:
    """统一名称中的全角字符及编号区间连接符。"""
    text = unicodedata.normalize("NFKC", clean_text(value)).upper()
    text = (
        text.replace("—", "-")
        .replace("–", "-")
        .replace("－", "-")
        .replace("～", "-")
        .replace("~", "-")
        .replace("至", "-")
    )
    text = re.sub(
        r"(\d+)\s*号?\s*-\s*(\d+)\s*号",
        lambda match: f"{match.group(1)}-{match.group(2)}号",
        text,
    )
    return text


def normalize_product_text(value: Any) -> str:
    """
    统一产品名称和PPT文件名。

    只删除需求中规定的通用字段；MYFUND、量化30、赤兔、稳健、均衡、
    积极、FOF等系列识别字段全部保留。
    """
    text = normalize_range_text(value)
    for part in sorted(GENERIC_PRODUCT_PARTS, key=len, reverse=True):
        normalized_part = unicodedata.normalize("NFKC", part).upper()
        text = text.replace(normalized_part, "")

    # 保留中文、英文字母、数字、“号”和区间连接符，其余符号全部删除。
    text = re.sub(r"[^0-9A-Z\u4e00-\u9fff号-]+", "", text)
    text = re.sub(r"-+", "-", text)
    return text.strip("-")


PRODUCT_RANGE_RE = re.compile(r"(\d+)\s*号?\s*-\s*(\d+)\s*号")
PRODUCT_NUMBER_RE = re.compile(r"(\d+)\s*号")


def parse_product_identity(product: ProductRecord) -> Optional[ProductIdentity]:
    """
    将产品清单中的单一产品拆分为“前缀＋编号＋后缀”。

    使用最后一个“数字＋号”作为产品编号，因此“量化30”中的30不会被误判。
    """
    normalized = normalize_product_text(product.name)
    if not normalized or PRODUCT_RANGE_RE.search(normalized):
        return None

    number_matches = list(PRODUCT_NUMBER_RE.finditer(normalized))
    if not number_matches:
        return None

    match = number_matches[-1]
    prefix = normalized[:match.start()]
    suffix = normalized[match.end():]
    if not prefix and not suffix:
        return None

    return ProductIdentity(
        product_name=product.name,
        prefix=prefix,
        suffix=suffix,
        number=int(match.group(1)),
        order=product.order,
    )


def strip_file_suffix_metadata(value: Any) -> str:
    """
    截取PPT文件名中的正式产品名称区段，并删除残余尾部元数据。

    仅用于文件名，不用于产品清单中的正式产品名称。匹配边界规则：
    - 文件名包含“万家共赢”时，删除其之前的事项编号、附件编号等内容；
    - 文件名包含“资产管理计划”时，删除该字段之后的全部内容；
    - 对没有“资产管理计划”的历史简称文件，继续删除末尾日期、版本号、
      流转备注和多余连接符。

    示例：
    “事项三：万家共赢中泰MYFUND均衡1号FOF集合资产管理计划-投决会上会材料”
    会被截取为：
    “万家共赢中泰MYFUND均衡1号FOF集合资产管理计划”。
    """
    text = unicodedata.normalize("NFKC", clean_text(value)).strip()

    # 文件扩展名通常已由Path.stem移除，这里保留防御性处理。
    text = re.sub(r"\.PPTX?$", "", text, flags=re.IGNORECASE)

    # 正式产品名以“万家共赢”作为优先起点。使用最后一次出现的位置，
    # 防止会议说明或文件夹拼接文本中更早出现同名公司字段。
    product_start = text.rfind("万家共赢")
    if product_start >= 0:
        text = text[product_start:]

    # “资产管理计划”是正式产品名称的终止边界。其后的“投决会上会材料”、
    # 日期、版本号、部门备注等均不参与产品匹配。
    product_end_marker = "资产管理计划"
    product_end = text.find(product_end_marker)
    if product_end >= 0:
        product_end += len(product_end_marker)
        text = text[:product_end]

    previous = None
    while text and text != previous:
        previous = text
        text = text.strip(" _-—–")

        # 常见复制/修订版本号：(1)、(1)(1)、（2）等。
        text = re.sub(r"(?:\s*[（(]\s*\d+\s*[）)])+$", "", text).strip()

        # 文件流转备注通常位于日期或版本号之后。
        text = re.sub(r"\s*[【\[][^】\]]{1,80}[】\]]$", "", text).strip()

        # 末尾日期：YYYYMMDD或YYYYMM。限定19/20开头，避免删除产品编号。
        text = re.sub(r"(?<!\d)(?:19|20)\d{6}$", "", text).strip()
        text = re.sub(r"(?<!\d)(?:19|20)\d{4}$", "", text).strip()

    return text.strip(" _-—–")


def parse_file_identity(value: Any) -> Optional[FileIdentity]:
    """从PPT文件名解析产品系列及单一编号或编号区间。"""
    normalized = normalize_product_text(strip_file_suffix_metadata(value))
    if not normalized:
        return None

    range_matches = list(PRODUCT_RANGE_RE.finditer(normalized))
    if range_matches:
        match = range_matches[-1]
        prefix = normalized[:match.start()]
        suffix = normalized[match.end():]
        if not prefix and not suffix:
            return None
        return FileIdentity(
            prefix=prefix,
            suffix=suffix,
            start_number=int(match.group(1)),
            end_number=int(match.group(2)),
            is_range=True,
        )

    number_matches = list(PRODUCT_NUMBER_RE.finditer(normalized))
    if not number_matches:
        return None

    match = number_matches[-1]
    prefix = normalized[:match.start()]
    suffix = normalized[match.end():]
    if not prefix and not suffix:
        return None
    number = int(match.group(1))
    return FileIdentity(
        prefix=prefix,
        suffix=suffix,
        start_number=number,
        end_number=number,
        is_range=False,
    )


def build_product_identities(
    products: Sequence[ProductRecord],
) -> dict[str, ProductIdentity]:
    """构建带“数字＋号”的产品身份；无编号产品留给第二遍匹配。"""
    identities: dict[str, ProductIdentity] = {}
    for product in products:
        normalized = normalize_product_text(product.name)
        if not normalized:
            LOGGER.warning(
                "产品识别", product.name, "N/A", "产品清单.xlsx",
                "产品名称标准化后为空，阶段二无法自动匹配PPT",
            )
            continue

        # 产品清单中的无编号产品不应被强行寻找“数字＋号”，
        # 它们会在文件第一遍未匹配后进入第二遍完整名称匹配。
        if not PRODUCT_NUMBER_RE.search(normalized):
            continue

        identity = parse_product_identity(product)
        if identity is None:
            LOGGER.warning(
                "产品识别", product.name, "N/A", "产品清单.xlsx",
                "产品名称包含编号字段，但无法拆分产品系列核心词与产品编号",
            )
            continue
        identities[product.name] = identity
    return identities


def build_numberless_product_identities(
    products: Sequence[ProductRecord],
) -> dict[str, NumberlessProductIdentity]:
    """构建产品名称中本身不含“数字＋号”的产品身份。"""
    identities: dict[str, NumberlessProductIdentity] = {}
    token_to_names: dict[str, list[str]] = {}

    for product in products:
        normalized = normalize_product_text(product.name)
        if not normalized or PRODUCT_NUMBER_RE.search(normalized):
            continue
        identity = NumberlessProductIdentity(
            product_name=product.name,
            full_token=normalized,
            order=product.order,
        )
        identities[product.name] = identity
        token_to_names.setdefault(normalized, []).append(product.name)

    for token, names in token_to_names.items():
        if len(names) > 1:
            LOGGER.warning(
                "产品识别", "、".join(names), "N/A", "产品清单.xlsx",
                f"多个无编号产品标准化后名称相同，第二遍文件匹配将要求人工确认（{token}）",
            )

    return identities


def identities_by_series(
    identities: dict[str, ProductIdentity],
) -> dict[str, list[ProductIdentity]]:
    grouped: dict[str, list[ProductIdentity]] = {}
    for identity in identities.values():
        grouped.setdefault(identity.series_core, []).append(identity)
    for values in grouped.values():
        values.sort(key=lambda item: (item.number, item.order, item.product_name))
    return grouped


def detect_product_scheme(product_name: str) -> str:
    normalized = compact_text(product_name)
    for canonical in SCHEME_DETECTION_ORDER:
        if canonical == "增强":
            # “指数增强/指增”属于指增方案，不应再次认定为增强方案。
            if "指增" in normalized or "指数增强" in normalized:
                continue
        if canonical in normalized:
            return canonical
    return ""


# ============================================================
# 5. Excel输入
# ============================================================


def require_path(path: Path, label: str) -> None:
    if not path.exists():
        LOGGER.error("文件扫描", "N/A", "N/A", str(path), f"未找到{label}")
        raise FileNotFoundError(f"未找到{label}：{path}")


def read_product_list(path: Path, header: str) -> list[ProductRecord]:
    require_path(path, "产品清单")
    dataframe = pd.read_excel(path, dtype=object, engine="openpyxl")
    normalized_columns = {compact_text(column): column for column in dataframe.columns}
    column = normalized_columns.get(compact_text(header))
    if column is None:
        raise ValueError(f"产品清单缺少表头：{header}")

    products: list[ProductRecord] = []
    seen: set[str] = set()
    for value in dataframe[column].tolist():
        name = clean_text(value)
        if not name:
            continue
        normalized = normalize_product_text(name)
        if not normalized or normalized in seen:
            if normalized in seen:
                LOGGER.warning(
                    "产品读取", name, "N/A", path.name,
                    "产品名称重复，已按单一产品处理",
                )
            continue
        seen.add(normalized)
        products.append(ProductRecord(name=name, order=len(products)))

    if not products:
        raise ValueError("产品清单中没有有效产品")

    LOGGER.success(
        "产品读取", "N/A", "N/A", path.name,
        f"共读取{len(products)}个待核查产品",
    )
    return products


def parse_plain_number(value: Any) -> Optional[float]:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return round(float(value), 4)
    text = clean_text(value).replace(",", "")
    if not text:
        return None
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return round(float(match.group(0)), 4)
    except ValueError:
        return None


def read_actual_strategies(path: Path, sheet_name: str) -> list[ActualStrategyRecord]:
    require_path(path, "阶段一输出文件")
    dataframe = pd.read_excel(path, sheet_name=sheet_name, dtype=object, engine="openpyxl")
    required = ("产品名称", "一级分类", "总净值占比")
    normalized_columns = {compact_text(column): column for column in dataframe.columns}
    missing = [name for name in required if compact_text(name) not in normalized_columns]
    if missing:
        raise ValueError(f"产品策略分类缺少表头：{'/'.join(missing)}")

    records: list[ActualStrategyRecord] = []
    for row in dataframe.itertuples(index=False, name=None):
        row_map = {column: value for column, value in zip(dataframe.columns, row)}
        product_name = clean_text(row_map[normalized_columns[compact_text("产品名称")]])
        category = clean_text(row_map[normalized_columns[compact_text("一级分类")]])
        ratio = parse_plain_number(row_map[normalized_columns[compact_text("总净值占比")]])
        if not product_name or not category or ratio is None:
            continue
        records.append(ActualStrategyRecord(product_name, category, ratio))

    LOGGER.success(
        "阶段一结果", "N/A", "N/A", path.name,
        f"读取{len(records)}条产品策略总净值占比",
    )
    return records


def read_secondary_strategies(
    path: Path,
    sheet_name: str,
) -> list[SecondaryStrategyRecord]:
    """
    从阶段一“科目比例明细”读取科目二级分类。

    使用“科目净值占比”逐科目汇总，避免重复使用每行都会出现的
    “产品策略总净值占比”。
    """
    require_path(path, "阶段一输出文件")
    dataframe = pd.read_excel(
        path,
        sheet_name=sheet_name,
        dtype=object,
        engine="openpyxl",
    )
    required = ("产品名称", "科目二级分类", "科目净值占比")
    normalized_columns = {
        compact_text(column): column for column in dataframe.columns
    }
    missing = [
        name for name in required
        if compact_text(name) not in normalized_columns
    ]
    if missing:
        raise ValueError(
            f"科目比例明细缺少表头：{'/'.join(missing)}"
        )

    records: list[SecondaryStrategyRecord] = []
    for row in dataframe.itertuples(index=False, name=None):
        row_map = {
            column: value
            for column, value in zip(dataframe.columns, row)
        }
        product_name = clean_text(
            row_map[normalized_columns[compact_text("产品名称")]]
        )
        secondary_category = clean_text(
            row_map[normalized_columns[compact_text("科目二级分类")]]
        )
        ratio = parse_plain_number(
            row_map[normalized_columns[compact_text("科目净值占比")]]
        )
        if not product_name or not secondary_category or ratio is None:
            continue
        records.append(
            SecondaryStrategyRecord(
                product_name=product_name,
                secondary_category=secondary_category,
                subject_ratio=ratio,
            )
        )

    LOGGER.success(
        "阶段一结果", "N/A", "二级分类", path.name,
        f"读取{len(records)}条科目二级分类净值占比",
    )
    return records


# ============================================================
# 6. PPT安全读取
# ============================================================


def iter_shapes_recursive(shapes: Iterable[Any]) -> Iterator[Any]:
    for shape in shapes:
        yield shape
        if isinstance(shape, GroupShape):
            try:
                yield from iter_shapes_recursive(shape.shapes)
            except (AttributeError, KeyError, NotImplementedError, TypeError, ValueError):
                continue


def slide_texts(slide: Any) -> list[str]:
    texts: list[str] = []
    try:
        shapes = slide.shapes
    except (AttributeError, KeyError, NotImplementedError, ValueError):
        return texts

    for shape in iter_shapes_recursive(shapes):
        try:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = clean_text(shape.text)
        except (AttributeError, KeyError, NotImplementedError, TypeError, ValueError):
            continue
        if text:
            texts.append(text)
    return texts




def slide_text_shape_records(slide: Any) -> list[TextShapeRecord]:
    """读取文本框文字及位置，供非原生表格配置识别。"""
    records: list[TextShapeRecord] = []
    try:
        shapes = slide.shapes
    except (AttributeError, KeyError, NotImplementedError, ValueError):
        return records

    for order, shape in enumerate(iter_shapes_recursive(shapes)):
        try:
            if not getattr(shape, "has_text_frame", False):
                continue
            value = clean_text(shape.text)
            if not value:
                continue
            records.append(
                TextShapeRecord(
                    text=value,
                    left=int(getattr(shape, "left", 0) or 0),
                    top=int(getattr(shape, "top", 0) or 0),
                    width=int(getattr(shape, "width", 0) or 0),
                    height=int(getattr(shape, "height", 0) or 0),
                    order=order,
                )
            )
        except (
            AttributeError,
            KeyError,
            NotImplementedError,
            TypeError,
            ValueError,
        ):
            continue
    return records
def slide_tables(slide: Any) -> Iterator[Any]:
    try:
        shapes = slide.shapes
    except (AttributeError, KeyError, NotImplementedError, ValueError):
        return

    for shape in iter_shapes_recursive(shapes):
        try:
            if getattr(shape, "has_table", False):
                yield shape.table
        except (AttributeError, KeyError, NotImplementedError, TypeError, ValueError):
            continue


def table_matrix(table: Any) -> list[list[str]]:
    matrix: list[list[str]] = []
    try:
        for row in table.rows:
            matrix.append([clean_text(cell.text) for cell in row.cells])
    except (AttributeError, KeyError, NotImplementedError, TypeError, ValueError):
        return []
    return matrix


def flatten_matrix(matrix: Sequence[Sequence[str]]) -> str:
    return "\n".join(cell for row in matrix for cell in row if clean_text(cell))


class PresentationReader:
    """
    逐文件打开、处理并释放PPT。

    V5.1.0起不再缓存Presentation对象或整份PPT的页面文本，避免处理大量PPT时
    内存随文件数量持续增长。正文匹配和策略提取都只保留当前文件。
    """

    @contextmanager
    def open(self, path: Path) -> Iterator[Optional[Presentation]]:
        presentation: Optional[Presentation] = None
        try:
            presentation = Presentation(str(path))
        except Exception as exc:
            LOGGER.error(
                "PPT读取", "N/A", "N/A", path.name,
                f"读取失败：{type(exc).__name__}: {exc}",
            )
            yield None
            return

        try:
            yield presentation
        finally:
            # python-pptx会在内存中展开XML关系、图片和图表对象。
            # 这里主动切断当前文件引用并触发垃圾回收，防止批量运行时持续累积。
            presentation = None
            gc.collect()

    @staticmethod
    def normalized_slide_text(slide: Any) -> str:
        parts = slide_texts(slide)
        for table in slide_tables(slide):
            table_text = flatten_matrix(table_matrix(table))
            if table_text:
                parts.append(table_text)
        return normalize_product_text("\n".join(parts))

    def find_first_match(
        self,
        path: Path,
        slide_limit: int,
        matcher: Callable[[str], PageMatchResult],
    ) -> tuple[Optional[int], Optional[PageMatchResult]]:
        """流式扫描页面，找到首个有效匹配后立即释放整份PPT。"""
        with self.open(path) as presentation:
            if presentation is None:
                return None, None

            try:
                for slide_number, slide in enumerate(presentation.slides, start=1):
                    if slide_limit > 0 and slide_number > slide_limit:
                        break
                    page_text = self.normalized_slide_text(slide)
                    if not page_text:
                        continue
                    result = matcher(page_text)
                    if result.ambiguous or result.product_names:
                        return slide_number, result
            except Exception as exc:
                LOGGER.error(
                    "PPT正文匹配", "N/A", "N/A", path.name,
                    f"扫描失败：{type(exc).__name__}: {exc}",
                )
                return None, None

        return None, None

    def extract_file(
        self,
        ppt_file: PptFile,
        product_name: str,
    ) -> tuple[Optional[PptMetadata], Optional[ExtractionResult], str]:
        """
        打开单个PPT并完成元数据与策略比例提取。

        返回：(元数据, 提取结果, 错误状态)。错误状态为空表示处理成功。
        """
        with self.open(ppt_file.path) as presentation:
            if presentation is None:
                return None, None, "PPT读取失败"

            try:
                metadata = extract_metadata(presentation)
                extraction = extract_allocation_plans(
                    presentation, ppt_file, product_name
                )
                return metadata, extraction, ""
            except Exception as exc:
                detail = f"PPT处理失败：{type(exc).__name__}: {exc}"
                LOGGER.error(
                    "PPT处理", product_name, "N/A", ppt_file.path.name, detail
                )
                return None, None, detail


# ============================================================
# 7. PPT扫描与产品匹配
# ============================================================


def scan_ppt_directory(directory: Path, file_type: str) -> list[PptFile]:
    require_path(directory, f"{file_type}目录")
    files = [
        PptFile(path=path, file_type=file_type)
        for path in directory.rglob("*")
        if path.is_file()
        and not path.name.startswith("~$")
        and path.suffix.lower() == ".pptx"
    ]
    files.sort(key=lambda item: str(item.path).lower())
    LOGGER.success(
        "文件扫描", "N/A", "N/A", directory.name,
        f"递归找到{len(files)}个PPTX文件，直接认定为{file_type}",
    )
    return files


def file_identity_is_incomplete(
    file_identity: FileIdentity,
    identities: dict[str, ProductIdentity],
) -> bool:
    """
    判断文件名解析出的系列核心词是否只是正式产品系列的一部分。

    例如“均衡3号推介材料”解析出的核心词只有“均衡”，不能据此认定为
    一个明确的其他产品系列，应继续扫描PPT封面或产品名称字段。
    """
    file_core = file_identity.series_core
    if not file_core:
        return True

    known_cores = {identity.series_core for identity in identities.values()}
    if file_core in known_cores:
        return False

    # 文件核心词是某个正式系列核心词的真子串，说明文件名信息不完整。
    return any(file_core in known_core for known_core in known_cores)


def match_file_identity_to_products(
    file_identity: FileIdentity,
    identities: dict[str, ProductIdentity],
) -> list[ProductIdentity]:
    """
    文件名严格匹配：产品系列核心词必须完全一致；随后比较单一编号或编号区间。
    """
    matched = [
        identity
        for identity in identities.values()
        if identity.series_core == file_identity.series_core
        and file_identity.contains(identity.number)
    ]
    return sorted(matched, key=lambda item: item.order)


def _range_pattern_for_identity(identity: ProductIdentity) -> re.Pattern[str]:
    return re.compile(
        re.escape(identity.prefix)
        + r"(\d+)\s*号?\s*-\s*(\d+)\s*号"
        + re.escape(identity.suffix)
    )


def match_products_in_page(
    page_text: str,
    identities: dict[str, ProductIdentity],
) -> PageMatchResult:
    """
    正文回退采用严格完整身份匹配。

    - 完整“前缀＋编号＋号＋后缀”可匹配单一产品；
    - 完整“前缀＋编号区间＋后缀”可匹配同系列多个产品；
    - 同一页命中不同产品系列时视为歧义；
    - 不生成“2号FOF”等宽泛尾部关键词。
    """
    normalized = normalize_product_text(page_text)
    if not normalized:
        return PageMatchResult()

    exact_matches: set[str] = set()
    range_matches: set[str] = set()
    range_details: set[str] = set()

    for identity in identities.values():
        if identity.full_token and identity.full_token in normalized:
            exact_matches.add(identity.product_name)

        pattern = _range_pattern_for_identity(identity)
        for match in pattern.finditer(normalized):
            start_number = int(match.group(1))
            end_number = int(match.group(2))
            low, high = sorted((start_number, end_number))
            if low <= identity.number <= high:
                range_matches.add(identity.product_name)
                range_details.add(f"{start_number}-{end_number}号")

    candidates = range_matches or exact_matches
    if not candidates:
        return PageMatchResult()

    series = {
        identities[name].series_core
        for name in candidates
        if name in identities
    }
    if len(series) > 1:
        return PageMatchResult(
            product_names=tuple(sorted(candidates)),
            ambiguous=True,
            detail="正文同页命中多个不同产品系列",
        )

    # 多个单一产品名称同时出现但没有区间表达，无法判断PPT实际归属。
    if not range_matches and len(exact_matches) > 1:
        return PageMatchResult(
            product_names=tuple(sorted(exact_matches)),
            ambiguous=True,
            detail="正文同页出现多个单一产品名称",
        )

    ordered = tuple(
        identity.product_name
        for identity in sorted(
            (identities[name] for name in candidates if name in identities),
            key=lambda item: item.order,
        )
    )
    detail = "、".join(sorted(range_details)) if range_matches else "完整产品名称"
    return PageMatchResult(
        product_names=ordered,
        is_range=bool(range_matches),
        ambiguous=False,
        detail=detail,
    )


def match_numberless_file_name(
    value: Any,
    identities: dict[str, NumberlessProductIdentity],
) -> list[NumberlessProductIdentity]:
    """第二遍文件名匹配：清理后完整名称必须与无编号产品完全一致。"""
    token = normalize_product_text(strip_file_suffix_metadata(value))
    if not token:
        return []
    matched = [
        identity
        for identity in identities.values()
        if identity.full_token == token
    ]
    return sorted(matched, key=lambda item: item.order)


def match_numberless_products_in_page(
    page_text: str,
    identities: dict[str, NumberlessProductIdentity],
) -> PageMatchResult:
    """
    第二遍正文匹配：只查找产品清单中本身没有“数字＋号”的完整产品名称。

    为避免短名称被较长名称包含造成误判，若同页命中名称存在包含关系，
    仅保留最长的完整名称；剩余多个不同产品时视为歧义。
    """
    normalized = normalize_product_text(page_text)
    if not normalized:
        return PageMatchResult()

    matched = [
        identity
        for identity in identities.values()
        if identity.full_token and identity.full_token in normalized
    ]
    if not matched:
        return PageMatchResult()

    longest_only = [
        identity
        for identity in matched
        if not any(
            identity.full_token != other.full_token
            and identity.full_token in other.full_token
            for other in matched
        )
    ]

    product_names = sorted(
        {identity.product_name for identity in longest_only},
        key=lambda name: identities[name].order,
    )
    if len(product_names) > 1:
        return PageMatchResult(
            product_names=tuple(product_names),
            ambiguous=True,
            detail="正文同页命中多个无编号产品完整名称",
        )

    return PageMatchResult(
        product_names=tuple(product_names),
        is_range=False,
        ambiguous=False,
        detail="无编号产品完整名称",
    )


def assign_products(
    assignments: dict[str, list[PptFile]],
    product_names: Sequence[str],
    ppt_file: PptFile,
) -> None:
    for product_name in product_names:
        assignments[product_name].append(ppt_file)


def assign_ppt_files(
    ppt_files: Sequence[PptFile],
    products: Sequence[ProductRecord],
    reader: PresentationReader,
    config: Config,
    stats: RunStats,
) -> dict[str, list[PptFile]]:
    """
    两遍文件匹配：

    第一遍：仅处理产品清单中带“数字＋号”的产品，使用系列核心词＋编号/区间。
    第二遍：仅处理第一遍未匹配的文件；对产品清单中本身无“数字＋号”的产品，
             使用清理后的完整产品名称严格匹配，不再强制寻找产品编号。

    对文件名完全无法表达产品归属的历史文件，第二遍完成无编号匹配后，
    仍可使用PPT正文对带号产品做原有的完整产品名称严格回退。
    """
    identities = build_product_identities(products)
    numberless_identities = build_numberless_product_identities(products)
    series_index = identities_by_series(identities)
    assignments = {product.name: [] for product in products}

    # (PPT文件, 第一遍未匹配原因, 是否允许第二遍继续按正文匹配带号产品)
    pending: list[tuple[PptFile, str, bool]] = []

    # --------------------------------------------------------
    # 第一遍：带号产品
    # --------------------------------------------------------
    for ppt_file in ppt_files:
        file_identity = parse_file_identity(ppt_file.path.stem)

        if file_identity is not None:
            matched_identities = match_file_identity_to_products(file_identity, identities)
            if matched_identities:
                assign_products(
                    assignments,
                    [identity.product_name for identity in matched_identities],
                    ppt_file,
                )
                stats.matched_files.add(ppt_file.path)

                for identity in matched_identities:
                    if file_identity.is_range:
                        LOGGER.success(
                            "文件匹配", identity.product_name, "N/A", ppt_file.path.name,
                            "第一遍编号区间匹配成功"
                            f"（产品核心词：{identity.series_core}；"
                            f"区间：{file_identity.range_label}；"
                            f"产品编号：{identity.number}号）",
                        )
                    else:
                        LOGGER.success(
                            "文件匹配", identity.product_name, "N/A", ppt_file.path.name,
                            "第一遍文件名严格匹配成功"
                            f"（产品核心词：{identity.series_core}；"
                            f"编号：{identity.number}号）",
                        )
                continue

            same_series = series_index.get(file_identity.series_core, [])
            if same_series:
                reason = (
                    "第一遍已识别带号产品系列，但编号不在产品清单或编号区间内"
                    f"（产品核心词：{file_identity.series_core}；"
                    f"文件编号：{file_identity.range_label}）"
                )
                pending.append((ppt_file, reason, False))
                LOGGER.info(
                    "文件匹配", "N/A", "N/A", ppt_file.path.name,
                    reason + "；进入第二遍无编号产品匹配",
                )
                continue

            if file_identity_is_incomplete(file_identity, identities):
                LOGGER.info(
                    "文件匹配", "N/A", "N/A", ppt_file.path.name,
                    "第一遍文件名仅包含不完整带号产品系列信息，"
                    "先执行PPT正文严格匹配"
                    f"（文件核心词：{file_identity.series_core}；"
                    f"文件编号：{file_identity.range_label}）",
                )

                if config.enable_content_matching:
                    matched_slide, content_result = reader.find_first_match(
                        ppt_file.path,
                        config.content_match_slide_limit,
                        lambda page_text: match_products_in_page(page_text, identities),
                    )

                    if content_result is not None and content_result.ambiguous:
                        reason = (
                            f"第一遍PPT正文第{matched_slide}页带号产品匹配存在歧义"
                            f"（{content_result.detail}）"
                        )
                        pending.append((ppt_file, reason, False))
                        LOGGER.warning(
                            "文件匹配", "N/A", "N/A", ppt_file.path.name,
                            reason + "；继续进入第二遍无编号产品匹配",
                        )
                        continue

                    if content_result is not None and content_result.product_names:
                        assign_products(assignments, content_result.product_names, ppt_file)
                        stats.matched_files.add(ppt_file.path)
                        for product_name in content_result.product_names:
                            identity = identities[product_name]
                            if content_result.is_range:
                                LOGGER.success(
                                    "文件匹配", product_name, "N/A", ppt_file.path.name,
                                    "第一遍PPT正文编号区间严格匹配成功"
                                    f"（第{matched_slide}页；产品核心词：{identity.series_core}；"
                                    f"区间：{content_result.detail}；产品编号：{identity.number}号）",
                                )
                            else:
                                LOGGER.success(
                                    "文件匹配", product_name, "N/A", ppt_file.path.name,
                                    "第一遍PPT正文完整产品名称匹配成功"
                                    f"（第{matched_slide}页；产品核心词：{identity.series_core}；"
                                    f"编号：{identity.number}号）",
                                )
                        continue

                reason = (
                    "第一遍不完整带号文件名及PPT正文均未匹配产品清单中的带号产品"
                )
                pending.append((ppt_file, reason, False))
                LOGGER.info(
                    "文件匹配", "N/A", "N/A", ppt_file.path.name,
                    reason + "；进入第二遍无编号产品匹配",
                )
                continue

            reason = (
                "第一遍文件名已识别为其他带号产品系列，"
                "为防止误配未执行带号产品正文回退"
                f"（文件产品核心词：{file_identity.series_core}；"
                f"文件编号：{file_identity.range_label}）"
            )
            pending.append((ppt_file, reason, False))
            LOGGER.info(
                "文件匹配", "N/A", "N/A", ppt_file.path.name,
                reason + "；进入第二遍无编号产品匹配",
            )
            continue

        # 文件名本身没有可解析的“数字＋号”。不要在这里强行寻找产品编号；
        # 先放入第二遍，优先匹配产品清单中本身无编号的产品。
        reason = "第一遍文件名未解析出产品编号"
        pending.append((ppt_file, reason, True))
        LOGGER.info(
            "文件匹配", "N/A", "N/A", ppt_file.path.name,
            reason + "；进入第二遍无编号产品匹配",
        )

    # --------------------------------------------------------
    # 第二遍：无编号产品；必要时再对无身份文件做带号正文回退
    # --------------------------------------------------------
    for ppt_file, first_pass_reason, allow_numbered_content_fallback in pending:
        matched_numberless = match_numberless_file_name(
            ppt_file.path.stem, numberless_identities
        )
        if len(matched_numberless) == 1:
            identity = matched_numberless[0]
            assign_products(assignments, [identity.product_name], ppt_file)
            stats.matched_files.add(ppt_file.path)
            LOGGER.success(
                "文件匹配", identity.product_name, "N/A", ppt_file.path.name,
                "第二遍无编号产品文件名完整匹配成功"
                f"（完整产品核心词：{identity.full_token}）",
            )
            continue

        if len(matched_numberless) > 1:
            stats.unmatched_files += 1
            stats.manual_review += 1
            LOGGER.warning(
                "文件匹配", "N/A", "N/A", ppt_file.path.name,
                "第二遍文件名同时匹配多个标准化后同名的无编号产品，"
                "需人工确认，已跳过",
            )
            continue

        # 文件名未完整命中时，可从PPT正文查找无编号产品的完整名称。
        numberless_content_result: Optional[PageMatchResult] = None
        numberless_slide: Optional[int] = None
        if config.enable_content_matching and numberless_identities:
            numberless_slide, numberless_content_result = reader.find_first_match(
                ppt_file.path,
                config.content_match_slide_limit,
                lambda page_text: match_numberless_products_in_page(
                    page_text, numberless_identities
                ),
            )

        if (
            numberless_content_result is not None
            and numberless_content_result.ambiguous
        ):
            stats.unmatched_files += 1
            stats.manual_review += 1
            LOGGER.warning(
                "文件匹配", "N/A", "N/A", ppt_file.path.name,
                f"第二遍PPT正文第{numberless_slide}页命中多个无编号产品，"
                "需人工确认，已跳过"
                f"（{numberless_content_result.detail}）",
            )
            continue

        if (
            numberless_content_result is not None
            and numberless_content_result.product_names
        ):
            assign_products(
                assignments, numberless_content_result.product_names, ppt_file
            )
            stats.matched_files.add(ppt_file.path)
            for product_name in numberless_content_result.product_names:
                identity = numberless_identities[product_name]
                LOGGER.success(
                    "文件匹配", product_name, "N/A", ppt_file.path.name,
                    "第二遍PPT正文无编号产品完整名称匹配成功"
                    f"（第{numberless_slide}页；完整产品核心词：{identity.full_token}）",
                )
            continue

        # 对文件名完全没有编号身份的历史文件，保留旧版正文回退能力：
        # 无编号产品未命中后，再按带号产品完整名称扫描正文。
        if (
            allow_numbered_content_fallback
            and config.enable_content_matching
            and identities
        ):
            numbered_content_result: Optional[PageMatchResult] = None
            numbered_slide: Optional[int] = None
            numbered_slide, numbered_content_result = reader.find_first_match(
                ppt_file.path,
                config.content_match_slide_limit,
                lambda page_text: match_products_in_page(page_text, identities),
            )

            if numbered_content_result is not None and numbered_content_result.ambiguous:
                stats.unmatched_files += 1
                stats.manual_review += 1
                LOGGER.warning(
                    "文件匹配", "N/A", "N/A", ppt_file.path.name,
                    f"第二遍PPT正文第{numbered_slide}页带号产品匹配存在歧义，"
                    "需人工确认，已跳过"
                    f"（{numbered_content_result.detail}）",
                )
                continue

            if numbered_content_result is not None and numbered_content_result.product_names:
                assign_products(assignments, numbered_content_result.product_names, ppt_file)
                stats.matched_files.add(ppt_file.path)
                for product_name in numbered_content_result.product_names:
                    identity = identities[product_name]
                    LOGGER.success(
                        "文件匹配", product_name, "N/A", ppt_file.path.name,
                        "第二遍PPT正文带号产品完整名称匹配成功"
                        f"（第{numbered_slide}页；产品核心词：{identity.series_core}；"
                        f"编号：{identity.number}号）",
                    )
                continue

        stats.unmatched_files += 1
        LOGGER.info(
            "文件匹配", "N/A", "N/A", ppt_file.path.name,
            "两遍匹配均未命中产品清单中的产品，已跳过"
            f"（第一遍原因：{first_pass_reason}）",
        )

    for files in assignments.values():
        # 同一路径在同一产品下只保留一次；区间文件仍可分配给多个产品。
        unique = {item.path: item for item in files}
        files[:] = sorted(unique.values(), key=lambda item: str(item.path).upper())

    return assignments


# ============================================================
# 8. 策略、方案与比例解析
# ============================================================


def match_strategy(value: Any) -> tuple[str, str]:
    """只从简短标签中识别策略，禁止返回整段公司或产品介绍。"""
    original = clean_text(value).replace("\n", " ").strip()
    normalized = compact_text(original)
    if not normalized or len(original) > 40:
        return "", ""

    candidates: list[tuple[int, str, str]] = []
    for category, aliases in STRATEGY_ALIASES:
        for alias in aliases:
            normalized_alias = compact_text(alias)
            if normalized_alias and normalized_alias in normalized:
                candidates.append((len(normalized_alias), alias, category))
    if not candidates:
        return "", ""

    _, _, category = max(candidates, key=lambda item: item[0])
    return original, category


def looks_like_unknown_strategy_label(value: Any) -> bool:
    text = clean_text(value).replace("\n", " ").strip()
    normalized = compact_text(text)
    if not normalized or len(text) > 30:
        return False
    if normalized in {compact_text(label) for label in GENERIC_TABLE_LABELS}:
        return False
    if match_scheme_label(text):
        return False
    if parse_ratio_range_checked(text)[2] != "未找到比例":
        return False
    if re.fullmatch(r"[\d.]+", normalized):
        return False
    return bool(re.search(r"[\u4e00-\u9fffA-Za-z]", text))


def match_scheme_label(value: Any) -> str:
    text = clean_text(value).replace("\n", " ").strip()
    normalized = compact_text(text)
    if not normalized or len(text) > 60:
        return ""

    # 策略单元格“指数增强”不能被误判为“增强方案”。
    _, strategy_category = match_strategy(text)

    for canonical in SCHEME_DETECTION_ORDER:
        aliases = SCHEME_ALIASES[canonical]
        for alias in aliases:
            normalized_alias = compact_text(alias)
            if normalized == normalized_alias:
                if canonical == "增强" and strategy_category == "指数增强":
                    continue
                return canonical
            if normalized_alias in normalized and any(marker in normalized for marker in ("型", "系列", "产品")):
                if canonical == "增强" and "指数增强" in normalized:
                    continue
                return canonical
    return ""


def allocation_table_context(matrix: Sequence[Sequence[str]]) -> bool:
    """仅使用四个指定关键词审核表头，不再接受其他比例/占比表述。"""
    valid_headers = tuple(
        compact_text(keyword)
        for keyword in ALLOCATION_HEADER_KEYWORDS
    )
    return any(
        any(header in compact_text(cell) for header in valid_headers)
        for row in matrix
        for cell in row
        if clean_text(cell)
    )


def _validate_ratio_bounds(
    lower: Optional[float],
    upper: Optional[float],
) -> tuple[Optional[float], Optional[float], str]:
    values = [value for value in (lower, upper) if value is not None]
    if any(value < 0 or value > 100 for value in values):
        return None, None, "策略比例异常"
    if lower is not None and upper is not None and lower > upper:
        return None, None, "策略比例异常"
    return lower, upper, "正常"


def parse_ratio_range_checked(
    value: Any,
) -> tuple[Optional[float], Optional[float], str]:
    text = clean_text(value)
    if not text:
        return None, None, "未找到比例"

    normalized = (
        text.replace("％", "%")
        .replace("—", "-")
        .replace("–", "-")
        .replace("~", "-")
        .replace("～", "-")
        .replace("至", "-")
        .replace("【", "")
        .replace("】", "")
        .replace("[", "")
        .replace("]", "")
        .replace("（", "")
        .replace("）", "")
        .replace("(", "")
        .replace(")", "")
        .replace(" ", "")
        .replace("\n", "")
    )

    # 必须明确含百分号，避免把年份、规模、注册资本等当作比例。
    if "%" not in normalized:
        return None, None, "未找到比例"

    interval_matches = list(re.finditer(
        r"(-?\d+(?:\.\d+)?)%?-(\-?\d+(?:\.\d+)?)%",
        normalized,
        flags=re.IGNORECASE,
    ))
    if len(interval_matches) > 1:
        intervals = {
            (round(float(match.group(1)), 4), round(float(match.group(2)), 4))
            for match in interval_matches
        }
        if len(intervals) > 1:
            return None, None, "策略比例异常"
    if interval_matches:
        match = interval_matches[0]
        return _validate_ratio_bounds(
            round(float(match.group(1)), 4),
            round(float(match.group(2)), 4),
        )

    upper_patterns = (
        r"不高于(-?\d+(?:\.\d+)?)%",
        r"不超过(-?\d+(?:\.\d+)?)%",
        r"至多(-?\d+(?:\.\d+)?)%",
        r"最高(-?\d+(?:\.\d+)?)%",
        r"(-?\d+(?:\.\d+)?)%以下",
        r"(-?\d+(?:\.\d+)?)%以内",
    )
    for pattern in upper_patterns:
        match = re.search(pattern, normalized)
        if match:
            return _validate_ratio_bounds(0.0, round(float(match.group(1)), 4))

    lower_patterns = (
        r"不低于(-?\d+(?:\.\d+)?)%",
        r"至少(-?\d+(?:\.\d+)?)%",
        r"(-?\d+(?:\.\d+)?)%以上",
    )
    for pattern in lower_patterns:
        match = re.search(pattern, normalized)
        if match:
            return _validate_ratio_bounds(round(float(match.group(1)), 4), 100.0)

    single = re.fullmatch(r"(-?\d+(?:\.\d+)?)%", normalized)
    if single:
        fixed = round(float(single.group(1)), 4)
        return _validate_ratio_bounds(fixed, fixed)

    # 单元格中含一个独立百分数，并伴随少量说明文字时仍可识别。
    singles = re.findall(r"(-?\d+(?:\.\d+)?)%", normalized)
    if len(singles) == 1:
        fixed = round(float(singles[0]), 4)
        return _validate_ratio_bounds(fixed, fixed)
    if len(set(singles)) > 1:
        return None, None, "策略比例异常"

    return None, None, "未找到比例"


def detect_strategy_header_columns(
    matrix: Sequence[Sequence[str]],
    header_rows: int = 4,
) -> dict[int, tuple[str, str]]:
    """识别方案位于行时的策略列，包括无法标准化的简短策略表头。"""
    columns: dict[int, tuple[str, str]] = {}
    if not matrix:
        return columns

    max_columns = max(len(row) for row in matrix)
    for column in range(max_columns):
        known: Optional[tuple[str, str]] = None
        unknown = ""
        for row_index in range(min(header_rows, len(matrix))):
            if column >= len(matrix[row_index]):
                continue
            cell = matrix[row_index][column]
            raw, category = match_strategy(cell)
            if category:
                known = (raw, category)
                break
            if looks_like_unknown_strategy_label(cell):
                unknown = clean_text(cell).replace("\n", " ")
        if known:
            columns[column] = known
        elif unknown:
            columns[column] = (unknown, "")
    return columns


def detect_scheme_cells(
    matrix: Sequence[Sequence[str]],
) -> tuple[list[tuple[int, int, str]], set[str]]:
    cells: list[tuple[int, int, str]] = []
    schemes: set[str] = set()
    for row_index, row in enumerate(matrix):
        for column_index, cell in enumerate(row):
            scheme = match_scheme_label(cell)
            if scheme:
                cells.append((row_index, column_index, clean_text(cell).replace("\n", " ")))
                schemes.add(scheme)
    return cells, schemes


def likely_strategy_label_in_row(
    row: Sequence[str],
    ratio_column: int,
) -> tuple[str, str]:
    # 优先已知策略。
    for cell in row:
        raw, category = match_strategy(cell)
        if category:
            return raw, category

    # 回退到比例列左侧的简短非通用文字。
    candidates = list(range(min(ratio_column, len(row)) - 1, -1, -1))
    candidates.extend(range(0, len(row)))
    seen: set[int] = set()
    for index in candidates:
        if index in seen or index >= len(row):
            continue
        seen.add(index)
        cell = row[index]
        if looks_like_unknown_strategy_label(cell):
            return clean_text(cell).replace("\n", " "), ""
    return "", ""


def make_plan(
    raw_strategy: str,
    category: str,
    ratio_text: Any,
    slide_number: int,
    scheme_label: str,
    source_order: int,
) -> Optional[ExtractedPlan]:
    lower, upper, ratio_status = parse_ratio_range_checked(ratio_text)
    if ratio_status == "未找到比例":
        return None
    status = "策略比例异常" if ratio_status == "策略比例异常" else "策略配置比例提取成功"
    return ExtractedPlan(
        raw_strategy=clean_text(raw_strategy).replace("\n", " ")[:80],
        primary_category=category,
        lower_limit=lower,
        upper_limit=upper,
        slide_number=slide_number,
        scheme_label=scheme_label,
        extraction_status=status,
        source_order=source_order,
    )


def extract_plans_from_table(
    matrix: Sequence[Sequence[str]],
    product_scheme: str,
    slide_number: int,
    table_order: int,
) -> tuple[list[ExtractedPlan], str]:
    """
    从一个配置表中提取计划。

    返回：(计划列表, 状态)
    状态可能为：成功、非配置表、配置方案未匹配、配置方案存在歧义、未提取到比例。
    """
    if not matrix or not allocation_table_context(matrix):
        return [], "非配置表"

    scheme_cells, schemes_present = detect_scheme_cells(matrix)
    target_cells = [
        cell for cell in scheme_cells
        if product_scheme and match_scheme_label(cell[2]) == product_scheme
    ]

    # 产品名称无法识别方案，但表中含多个不同方案时，不猜测。
    if not product_scheme and len(schemes_present) > 1:
        return [], "配置方案存在歧义"

    strategy_columns = detect_strategy_header_columns(matrix)
    plans: list[ExtractedPlan] = []
    sequence = table_order * 10000

    # A. 方案位于行：目标方案行 + 顶部策略列。
    if target_cells and strategy_columns:
        target_rows = sorted({row_index for row_index, _, _ in target_cells})
        for row_index in target_rows:
            row = matrix[row_index]
            scheme_labels = [
                label for r, _, label in target_cells if r == row_index
            ]
            scheme_label = " / ".join(dict.fromkeys(scheme_labels))
            for column_index, (raw_strategy, category) in strategy_columns.items():
                if column_index >= len(row):
                    continue
                sequence += 1
                plan = make_plan(
                    raw_strategy, category, row[column_index], slide_number,
                    scheme_label, sequence,
                )
                if plan:
                    plans.append(plan)
        if plans:
            return plans, "成功"

    # B. 方案位于列：目标方案列 + 各策略行。
    if target_cells:
        target_columns = sorted({column_index for _, column_index, _ in target_cells})
        for column_index in target_columns:
            labels = [label for _, c, label in target_cells if c == column_index]
            scheme_label = " / ".join(dict.fromkeys(labels))
            for row in matrix:
                if column_index >= len(row):
                    continue
                raw_strategy, category = likely_strategy_label_in_row(row, column_index)
                if not raw_strategy:
                    continue
                sequence += 1
                plan = make_plan(
                    raw_strategy, category, row[column_index], slide_number,
                    scheme_label, sequence,
                )
                if plan:
                    plans.append(plan)
        if plans:
            return plans, "成功"

    # C. 表内出现方案，但没有目标方案。
    if product_scheme and schemes_present and not target_cells:
        return [], "配置方案未匹配"

    # D. 表格已通过严格表头审核后，识别真正存放百分比数据的列。
    # 审核表头只负责确认该表属于策略配置表；实际比例列仍按百分比数据定位。
    max_columns = max(len(row) for row in matrix)

    def ratio_value_count(column_index: int) -> int:
        return sum(
            1
            for row in matrix
            if column_index < len(row)
            and parse_ratio_range_checked(row[column_index])[2]
            in {"正常", "策略比例异常"}
        )

    ratio_columns: set[int] = {
        column_index
        for column_index in range(max_columns)
        if ratio_value_count(column_index) > 0
    }

    if len(ratio_columns) > 1 and not product_scheme:
        return [], "配置方案存在歧义"

    for column_index in sorted(ratio_columns):
        for row in matrix:
            if column_index >= len(row):
                continue
            raw_strategy, category = likely_strategy_label_in_row(row, column_index)
            if not raw_strategy:
                continue
            sequence += 1
            plan = make_plan(
                raw_strategy, category, row[column_index], slide_number,
                "唯一配置方案", sequence,
            )
            if plan:
                plans.append(plan)

    return (plans, "成功") if plans else ([], "未提取到比例")


def deduplicate_plans(plans: Sequence[ExtractedPlan]) -> list[ExtractedPlan]:
    """同一配置方案内部去重；不同配置方案即使比例相同也全部保留。"""
    seen: set[tuple[str, str, Optional[float], Optional[float], int, str]] = set()
    result: list[ExtractedPlan] = []
    for plan in sorted(plans, key=lambda item: item.source_order):
        key = (
            plan.scheme_label,
            plan.raw_strategy,
            plan.lower_limit,
            plan.upper_limit,
            plan.slide_number,
            plan.extraction_status,
        )
        if key in seen:
            continue
        seen.add(key)
        result.append(plan)
    return result


def is_promotion_allocation_page(
    slide_content: str,
    product_scheme: str,
) -> bool:
    """
    判断宣推PPT页面是否属于策略配置页面。

    除“投资方案/资产配置方案”外，兼容与目标产品方案一致的
    “积极型策略组合”“均衡型策略组合”等历史模板。只有页面同时出现
    目标方案标签与“策略组合”类标题时才放行，避免扫描普通市场分析页。
    """
    normalized = compact_text(slide_content)
    if any(compact_text(keyword) in normalized for keyword in PROMOTION_PAGE_KEYWORDS):
        return True

    if not product_scheme:
        return False

    scheme_hit = any(
        compact_text(alias) in normalized
        for alias in SCHEME_ALIASES.get(product_scheme, ())
    )
    combination_hit = any(
        compact_text(keyword) in normalized
        for keyword in PROMOTION_SCHEME_PAGE_KEYWORDS
    )
    return scheme_hit and combination_hit




def is_text_allocation_candidate_page(
    slide_content: str,
    product_scheme: str,
) -> bool:
    """
    文本框补充识别只在配置语境页面启用，避免把市场回顾页收益率误判为策略比例。
    """
    normalized = compact_text(slide_content)
    if any(
        compact_text(keyword) in normalized
        for keyword in (*TEXT_ALLOCATION_PAGE_KEYWORDS, *ALLOCATION_HEADER_KEYWORDS)
    ):
        return True

    if product_scheme:
        scheme_hit = any(
            compact_text(alias) in normalized
            for alias in SCHEME_ALIASES.get(product_scheme, ())
        )
        if scheme_hit and ("策略" in normalized or "配置" in normalized):
            return True
    return False


def _strict_text_ratio(
    value: Any,
) -> tuple[Optional[float], Optional[float], str]:
    """文本框模式只接受一个完整的 x%-y% 区间。"""
    text = clean_text(value)
    matches = list(TEXT_RATIO_PATTERN.finditer(text))
    if not matches:
        return None, None, "未找到比例"

    intervals = {
        (
            round(float(match.group("lower")), 4),
            round(float(match.group("upper")), 4),
        )
        for match in matches
    }
    if len(intervals) != 1:
        return None, None, "策略比例异常"

    lower, upper = next(iter(intervals))
    return _validate_ratio_bounds(lower, upper)


def _clean_text_strategy_prefix(value: Any) -> str:
    text = clean_text(value).replace("\n", " ").strip()
    # “###”代表冒号、点线、横线、井号、空格等视觉分隔。
    text = re.sub(r"[\s:：|｜#＃·•.…\-—–_]+$", "", text).strip()
    text = re.sub(r"^[\s:：|｜#＃·•.…\-—–_]+", "", text).strip()
    return text


def _strategy_from_text_candidate(value: Any) -> tuple[str, str]:
    candidate = _clean_text_strategy_prefix(value)
    if not candidate or len(candidate) > 40:
        return "", ""
    if compact_text(candidate) in {
        compact_text(label) for label in TEXT_NON_STRATEGY_LABELS
    }:
        return "", ""
    raw, category = match_strategy(candidate)
    if raw:
        return raw, category
    if looks_like_unknown_strategy_label(candidate):
        return candidate, ""
    return "", ""


def _plans_inside_text_shape(
    record: TextShapeRecord,
    slide_number: int,
    sequence_start: int,
) -> list[ExtractedPlan]:
    """识别同一文本框内“策略名称 ### x%-y%”及上下相邻行。"""
    plans: list[ExtractedPlan] = []
    sequence = sequence_start
    lines = [clean_text(line) for line in record.text.splitlines() if clean_text(line)]

    # 同一行：策略名称 ### x%-y%
    for line in lines:
        for match in TEXT_RATIO_PATTERN.finditer(line):
            prefix = line[:match.start()]
            raw, category = _strategy_from_text_candidate(prefix)
            if not raw:
                continue
            lower, upper, status = _strict_text_ratio(match.group(0))
            if status == "未找到比例":
                continue
            sequence += 1
            plans.append(
                ExtractedPlan(
                    raw_strategy=raw[:80],
                    primary_category=category,
                    lower_limit=lower,
                    upper_limit=upper,
                    slide_number=slide_number,
                    scheme_label="文本框补充识别",
                    extraction_status=(
                        "策略比例异常"
                        if status == "策略比例异常"
                        else "策略配置比例提取成功"
                    ),
                    source_order=sequence,
                )
            )

    # 上一行是策略、下一行仅为比例。
    for line_index, line in enumerate(lines):
        stripped = line.strip()
        ratio_match = TEXT_RATIO_PATTERN.fullmatch(stripped)
        if ratio_match is None or line_index == 0:
            continue
        raw, category = _strategy_from_text_candidate(lines[line_index - 1])
        if not raw:
            continue
        lower, upper, status = _strict_text_ratio(stripped)
        if status == "未找到比例":
            continue
        sequence += 1
        plans.append(
            ExtractedPlan(
                raw_strategy=raw[:80],
                primary_category=category,
                lower_limit=lower,
                upper_limit=upper,
                slide_number=slide_number,
                scheme_label="文本框补充识别",
                extraction_status=(
                    "策略比例异常"
                    if status == "策略比例异常"
                    else "策略配置比例提取成功"
                ),
                source_order=sequence,
            )
        )

    return plans


def _horizontal_overlap(first: TextShapeRecord, second: TextShapeRecord) -> int:
    return max(0, min(first.right, second.right) - max(first.left, second.left))


def _pair_text_shape_score(
    strategy_record: TextShapeRecord,
    ratio_record: TextShapeRecord,
    slide_width: int,
    slide_height: int,
    known_strategy: bool,
) -> Optional[float]:
    """
    名称与比例分处两个文本框时按空间关系打分：
    - 同行且名称在左侧优先；
    - 同列且名称在上方次之；
    - 关键词库已识别的策略加分。
    """
    width = max(slide_width, 1)
    height = max(slide_height, 1)
    known_bonus = 30.0 if known_strategy else 0.0

    # 同行：垂直中心接近，策略框位于比例框左侧。
    vertical_distance = abs(strategy_record.center_y - ratio_record.center_y)
    same_row_tolerance = max(
        strategy_record.height,
        ratio_record.height,
        int(height * 0.035),
    )
    horizontal_gap = ratio_record.left - strategy_record.right
    if (
        strategy_record.center_x < ratio_record.center_x
        and vertical_distance <= same_row_tolerance
        and horizontal_gap <= width * 0.35
        and horizontal_gap >= -width * 0.05
    ):
        return (
            80.0
            + known_bonus
            - 30.0 * vertical_distance / max(same_row_tolerance, 1)
            - 20.0 * max(horizontal_gap, 0) / max(width * 0.35, 1)
        )

    # 同列：策略框位于比例框上方，横向有较大重叠。
    overlap = _horizontal_overlap(strategy_record, ratio_record)
    min_width = max(min(strategy_record.width, ratio_record.width), 1)
    overlap_ratio = overlap / min_width
    vertical_gap = ratio_record.top - strategy_record.bottom
    if (
        strategy_record.center_y < ratio_record.center_y
        and overlap_ratio >= 0.30
        and vertical_gap <= height * 0.18
        and vertical_gap >= -height * 0.03
    ):
        return (
            60.0
            + known_bonus
            + 15.0 * overlap_ratio
            - 20.0 * max(vertical_gap, 0) / max(height * 0.18, 1)
        )

    return None


def extract_plans_from_text_shapes(
    slide: Any,
    slide_number: int,
    slide_width: int,
    slide_height: int,
) -> list[ExtractedPlan]:
    """
    从配置页面文本框补充提取策略比例。

    第一优先级：同一文本框内“策略名称 ### x%-y%”；
    第二优先级：仅含策略名的文本框与仅含比例的相邻文本框配对。
    """
    records = slide_text_shape_records(slide)
    plans: list[ExtractedPlan] = []
    sequence_base = 900000 + slide_number * 1000

    for record in records:
        plans.extend(
            _plans_inside_text_shape(
                record,
                slide_number,
                sequence_base + record.order * 10,
            )
        )

    strategy_records: list[tuple[TextShapeRecord, str, str]] = []
    ratio_records: list[tuple[TextShapeRecord, str]] = []

    for record in records:
        one_line = clean_text(record.text).replace("\n", " ").strip()
        ratio_match = TEXT_RATIO_PATTERN.fullmatch(one_line)
        if ratio_match:
            ratio_records.append((record, ratio_match.group(0)))
            continue

        raw, category = _strategy_from_text_candidate(one_line)
        if raw:
            strategy_records.append((record, raw, category))

    for ratio_index, (ratio_record, ratio_text) in enumerate(ratio_records):
        scored: list[
            tuple[float, TextShapeRecord, str, str]
        ] = []
        for strategy_record, raw, category in strategy_records:
            score = _pair_text_shape_score(
                strategy_record,
                ratio_record,
                slide_width,
                slide_height,
                bool(category),
            )
            if score is not None:
                scored.append((score, strategy_record, raw, category))

        if not scored:
            continue
        scored.sort(key=lambda item: (-item[0], item[1].order))
        best = scored[0]

        # 前两名过于接近时不自动配对，防止同页多个策略框错配。
        if len(scored) > 1 and best[0] - scored[1][0] < 12:
            continue

        lower, upper, status = _strict_text_ratio(ratio_text)
        if status == "未找到比例":
            continue
        plans.append(
            ExtractedPlan(
                raw_strategy=best[2][:80],
                primary_category=best[3],
                lower_limit=lower,
                upper_limit=upper,
                slide_number=slide_number,
                scheme_label="文本框补充识别",
                extraction_status=(
                    "策略比例异常"
                    if status == "策略比例异常"
                    else "策略配置比例提取成功"
                ),
                source_order=sequence_base + 500 + ratio_index,
            )
        )

    return deduplicate_plans(plans)
def extract_allocation_plans(
    presentation: Presentation,
    ppt_file: PptFile,
    product_name: str,
) -> ExtractionResult:
    product_scheme = detect_product_scheme(product_name)
    candidate_slides: list[tuple[int, Any]] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        if ppt_file.file_type == "宣推文件":
            # V5.2.1：宣推文件不再先查找“资产配置方案/投资方案”页面。
            # 直接将全部页面交给原生表格扫描；表格未提取到策略时，
            # 后续再按原有文本框补充识别逻辑处理。
            candidate_slides.append((slide_number, slide))
            continue

        # 投决会逻辑保持不变：全部页面直接进入原生表格扫描。
        candidate_slides.append((slide_number, slide))

    # 第一遍仍按原有原生表格逻辑处理。
    plans: list[ExtractedPlan] = []
    table_seen = False
    scheme_unmatched = False
    scheme_ambiguous = False

    table_order = 0
    for slide_number, slide in candidate_slides:
        for table in slide_tables(slide):
            table_order += 1
            matrix = table_matrix(table)
            if not allocation_table_context(matrix):
                continue
            table_seen = True
            table_plans, table_status = extract_plans_from_table(
                matrix,
                product_scheme,
                slide_number,
                table_order,
            )
            plans.extend(table_plans)
            if table_status == "配置方案未匹配":
                scheme_unmatched = True
            elif table_status == "配置方案存在歧义":
                scheme_ambiguous = True

    plans = deduplicate_plans(plans)
    if plans:
        return ExtractionResult(tuple(plans), "成功")

    # 第二遍仅在原生表格未提取到任何策略时启用文本框补充识别。
    text_plans: list[ExtractedPlan] = []
    slide_width = int(getattr(presentation, "slide_width", 0) or 0)
    slide_height = int(getattr(presentation, "slide_height", 0) or 0)

    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_content = "\n".join(slide_texts(slide))
        slide_content += "\n" + "\n".join(
            flatten_matrix(table_matrix(table))
            for table in slide_tables(slide)
        )
        if not is_text_allocation_candidate_page(
            slide_content,
            product_scheme,
        ):
            continue

        text_plans.extend(
            extract_plans_from_text_shapes(
                slide,
                slide_number,
                slide_width,
                slide_height,
            )
        )

    text_plans = deduplicate_plans(text_plans)
    if text_plans:
        return ExtractionResult(
            tuple(text_plans),
            "文本框补充识别成功",
        )

    if not candidate_slides:
        return ExtractionResult(
            (),
            "未找到资产配置方案/投资方案页面",
        )
    if scheme_unmatched:
        return ExtractionResult(
            (),
            "配置方案未匹配，需人工确认",
        )
    if scheme_ambiguous:
        return ExtractionResult(
            (),
            "配置方案存在歧义，需人工确认",
        )
    if table_seen:
        return ExtractionResult((), "未找到策略配置比例")
    return ExtractionResult((), "未找到策略配置比例")


# ============================================================
# 9. 产品要素字段提取
# ============================================================


def _all_metadata_labels() -> list[str]:
    return sorted(
        {
            *TARGET_CUSTOMER_HEADERS,
            *(alias for aliases in METADATA_FIELDS.values() for alias in aliases),
        },
        key=len,
        reverse=True,
    )


def extract_labeled_value_from_text(text: str, aliases: Sequence[str]) -> str:
    """
    从文本块拆分普通产品要素字段。

    普通文本要求字段后存在“：/为/是”，避免从风险揭示长段落误取；
    当同一文本块中同时出现至少两个产品要素标签时，允许省略分隔符。
    “目标客户”和“风险收益特征”不走此文本框回退，只从原生表格中按表头右侧提取。
    """
    cleaned = clean_text(text).replace("\n", " ")
    all_labels = _all_metadata_labels()
    normalized = compact_text(cleaned)
    label_hit_count = sum(
        1 for label in all_labels if compact_text(label) in normalized
    )
    delimiter = r"(?:\s*[：:]\s*|\s*(?:为|是)\s*|\s*)" if label_hit_count >= 2 else r"(?:\s*[：:]\s*|\s*(?:为|是)\s*)"
    next_label = "|".join(re.escape(label) for label in all_labels)

    for alias in sorted(aliases, key=len, reverse=True):
        pattern = re.compile(
            rf"{re.escape(alias)}{delimiter}(.*?)"
            rf"(?=(?:{next_label})(?:\s*[：:]|\s*(?:为|是))?|$)",
            re.IGNORECASE,
        )
        match = pattern.search(cleaned)
        if match:
            value = clean_text(match.group(1)).strip("；;，,。 ")
            if value:
                return value[:300]
    return ""


def is_metadata_label(value: Any) -> bool:
    normalized = compact_text(value).strip("：:")
    if normalized in {compact_text(header) for header in TARGET_CUSTOMER_HEADERS}:
        return True
    return any(
        compact_text(alias) in normalized
        for aliases in METADATA_FIELDS.values()
        for alias in aliases
    )


def normalize_metadata_field_value(field_name: str, value: Any) -> str:
    """清洗普通产品要素字段值。"""
    del field_name
    return clean_text(value)[:300]


def field_from_matrix(
    matrix: Sequence[Sequence[str]],
    field_name: str,
    aliases: Sequence[str],
) -> str:
    normalized_aliases = [compact_text(alias) for alias in aliases]
    for row_index, row in enumerate(matrix):
        for column_index, cell in enumerate(row):
            normalized_cell = compact_text(cell)
            if not any(alias in normalized_cell for alias in normalized_aliases):
                continue

            # 同格中可能同时存在多个字段。
            same_cell = extract_labeled_value_from_text(cell, aliases)
            normalized_value = normalize_metadata_field_value(field_name, same_cell)
            if normalized_value:
                return normalized_value

            # 键值表：字段在左，值在右。
            if column_index + 1 < len(row):
                right = clean_text(row[column_index + 1])
                if right and not is_metadata_label(right):
                    normalized_value = normalize_metadata_field_value(field_name, right)
                    if normalized_value:
                        return normalized_value

            # 横向表头：字段在上，值在下。
            if row_index + 1 < len(matrix) and column_index < len(matrix[row_index + 1]):
                below = clean_text(matrix[row_index + 1][column_index])
                if below and not is_metadata_label(below):
                    normalized_value = normalize_metadata_field_value(field_name, below)
                    if normalized_value:
                        return normalized_value
    return ""


def target_customer_from_matrix(matrix: Sequence[Sequence[str]]) -> str:
    """
    严格查找原生表格中的“目标客户”或“风险收益特征”表头，
    并抄写同一行右侧全部文字。

    - 表头单元格标准化后必须严格等于其中一个关键词；
    - 不解析目标收益、最大回撤，也不限制文字格式；
    - 右侧所有非空单元格按原顺序以换行连接；
    - 不截断内容。
    """
    expected_headers = {compact_text(header) for header in TARGET_CUSTOMER_HEADERS}
    for row in matrix:
        for column_index, cell in enumerate(row):
            normalized_cell = compact_text(cell).strip("：:")
            if normalized_cell not in expected_headers:
                continue

            right_values = [
                clean_text(value)
                for value in row[column_index + 1:]
                if clean_text(value)
            ]
            if right_values:
                return "\n".join(right_values)
    return ""


def extract_metadata(presentation: Presentation) -> PptMetadata:
    values = {
        "target_customer": "",
        "product_positioning": "",
    }

    candidate_slides: list[tuple[int, int, Any]] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = slide_texts(slide)
        table_texts = [flatten_matrix(table_matrix(table)) for table in slide_tables(slide)]
        all_text = "\n".join([*texts, *table_texts])
        normalized = compact_text(all_text)
        if compact_text("产品要素") in normalized:
            priority = 0
        elif compact_text("风险收益特征") in normalized:
            priority = 1
        elif any(compact_text(keyword) in normalized for keyword in METADATA_AUXILIARY_PAGE_KEYWORDS):
            priority = 2
        else:
            continue
        candidate_slides.append((priority, slide_number, slide))

    candidate_slides.sort(key=lambda item: (item[0], item[1]))

    for _, _, slide in candidate_slides:
        tables = list(slide_tables(slide))
        # 原生表格优先。“目标客户”和“风险收益特征”只从表格中读取。
        for table in tables:
            matrix = table_matrix(table)

            if not values["target_customer"]:
                target_customer = target_customer_from_matrix(matrix)
                if target_customer:
                    values["target_customer"] = target_customer

            if not values["product_positioning"]:
                product_positioning = field_from_matrix(
                    matrix,
                    "product_positioning",
                    METADATA_FIELDS["product_positioning"],
                )
                if product_positioning:
                    values["product_positioning"] = product_positioning

        # 产品定位保留文本框回退；目标客户相关关键词不做文本框回退。
        if not values["product_positioning"]:
            for text in slide_texts(slide):
                raw_value = extract_labeled_value_from_text(
                    text,
                    METADATA_FIELDS["product_positioning"],
                )
                value = normalize_metadata_field_value(
                    "product_positioning",
                    raw_value,
                )
                if value:
                    values["product_positioning"] = value
                    break

        if all(values.values()):
            break

    return PptMetadata(
        target_customer=values["target_customer"],
        product_positioning=values["product_positioning"],
    )


# ============================================================
# 10. 核查与行生成
# ============================================================


def compare_ratio(
    actual: Optional[float],
    lower: Optional[float],
    upper: Optional[float],
) -> str:
    if actual is None:
        return "未找到对应产品策略"
    if lower is None or upper is None:
        return "计划比例缺失"
    if actual < lower:
        return "低于计划下限"
    if actual > upper:
        return "高于计划上限"
    return "符合"


def build_actual_lookup(
    records: Sequence[ActualStrategyRecord],
    products: Sequence[ProductRecord],
) -> dict[tuple[str, str], float]:
    canonical_products = {
        normalize_product_text(product.name): product.name for product in products
    }
    lookup: dict[tuple[str, str], float] = {}
    for record in records:
        product_name = canonical_products.get(normalize_product_text(record.product_name))
        if not product_name:
            continue
        key = (product_name, record.primary_category)
        lookup[key] = round(lookup.get(key, 0.0) + record.total_ratio, 4)
    return lookup




def build_secondary_lookup(
    records: Sequence[SecondaryStrategyRecord],
    products: Sequence[ProductRecord],
) -> dict[str, tuple[SecondaryStrategyRecord, ...]]:
    canonical_products = {
        normalize_product_text(product.name): product.name
        for product in products
    }
    grouped: dict[str, list[SecondaryStrategyRecord]] = {}

    for record in records:
        product_name = canonical_products.get(
            normalize_product_text(record.product_name)
        )
        if not product_name:
            continue
        grouped.setdefault(product_name, []).append(
            SecondaryStrategyRecord(
                product_name=product_name,
                secondary_category=record.secondary_category,
                subject_ratio=record.subject_ratio,
            )
        )

    return {
        product_name: tuple(values)
        for product_name, values in grouped.items()
    }


def secondary_actual_for_file_strategy(
    product_name: str,
    file_strategy: str,
    secondary_lookup: dict[
        str,
        tuple[SecondaryStrategyRecord, ...],
    ],
) -> tuple[Optional[float], str]:
    """
    一级分类未找到时，按文件策略文字中的“可转债/宏观”汇总科目二级分类。

    同时出现两个关键词时，按科目行去重后合计全部可转债与宏观占比。
    """
    normalized_strategy = compact_text(file_strategy)
    requested = [
        keyword
        for keyword in SECONDARY_FALLBACK_KEYWORDS
        if compact_text(keyword) in normalized_strategy
    ]
    if not requested:
        return None, ""

    total = 0.0
    matched = False
    for record in secondary_lookup.get(product_name, ()):
        normalized_category = compact_text(record.secondary_category)
        if any(
            compact_text(keyword) in normalized_category
            for keyword in requested
        ):
            total += record.subject_ratio
            matched = True

    if not matched:
        return None, "、".join(requested)

    return round(total, 4), "、".join(requested)


def build_check_rows(
    products: Sequence[ProductRecord],
    assignments: dict[str, list[PptFile]],
    actual_lookup: dict[tuple[str, str], float],
    secondary_lookup: dict[
        str,
        tuple[SecondaryStrategyRecord, ...],
    ],
    reader: PresentationReader,
    stats: RunStats,
) -> list[CheckRow]:
    rows: list[CheckRow] = []
    type_order = {"投决会文件": 0, "宣推文件": 1}

    for product in products:
        product_files = assignments.get(product.name, [])
        files_by_type = {
            file_type: sorted(
                [item for item in product_files if item.file_type == file_type],
                key=lambda item: str(item.path).lower(),
            )
            for file_type in ("投决会文件", "宣推文件")
        }

        for file_type in ("投决会文件", "宣推文件"):
            files = files_by_type[file_type]
            if not files:
                if file_type == "投决会文件":
                    stats.products_without_committee += 1
                else:
                    stats.products_without_promotion += 1
                LOGGER.error(
                    "文件匹配", product.name, "N/A", "N/A",
                    f"未找到{file_type}",
                )
                rows.append(CheckRow(
                    product_name=product.name,
                    primary_category="",
                    file_name="N/A",
                    file_type=file_type,
                    file_strategy="",
                    actual_total_ratio=None,
                    lower_limit=None,
                    upper_limit=None,
                    check_result="未找到相关文件",
                    product_order=product.order,
                ))
                continue

            for ppt_file in files:
                metadata, extraction, processing_error = reader.extract_file(
                    ppt_file, product.name
                )
                if processing_error:
                    rows.append(CheckRow(
                        product_name=product.name,
                        primary_category="",
                        file_name=ppt_file.path.name,
                        file_type=file_type,
                        file_strategy="",
                        actual_total_ratio=None,
                        lower_limit=None,
                        upper_limit=None,
                        check_result=processing_error,
                        product_order=product.order,
                    ))
                    continue

                if metadata is None or extraction is None:
                    rows.append(CheckRow(
                        product_name=product.name,
                        primary_category="",
                        file_name=ppt_file.path.name,
                        file_type=file_type,
                        file_strategy="",
                        actual_total_ratio=None,
                        lower_limit=None,
                        upper_limit=None,
                        check_result="PPT处理失败：未返回提取结果",
                        product_order=product.order,
                    ))
                    continue

                if not extraction.plans:
                    status = extraction.status
                    if "需人工确认" in status:
                        stats.manual_review += 1
                        LOGGER.warning(
                            "PPT提取", product.name, "N/A", ppt_file.path.name, status
                        )
                    else:
                        LOGGER.error(
                            "PPT提取", product.name, "N/A", ppt_file.path.name, status
                        )
                    rows.append(CheckRow(
                        product_name=product.name,
                        primary_category="",
                        file_name=ppt_file.path.name,
                        file_type=file_type,
                        file_strategy="",
                        actual_total_ratio=None,
                        lower_limit=None,
                        upper_limit=None,
                        check_result=status,
                        target_customer=metadata.target_customer,
                        product_positioning=metadata.product_positioning,
                        product_order=product.order,
                    ))
                    continue

                LOGGER.success(
                    "PPT提取", product.name, "策略", ppt_file.path.name,
                    f"提取{len(extraction.plans)}条策略配置比例",
                )

                if extraction.status == "文本框补充识别成功":
                    stats.text_fallback_plans += len(extraction.plans)

                for plan in extraction.plans:
                    stats.extracted_plans += 1
                    resolved_category = plan.primary_category
                    actual = (
                        actual_lookup.get(
                            (product.name, plan.primary_category)
                        )
                        if plan.primary_category
                        else None
                    )

                    # 一级分类没有对应策略时，再按文件策略中的“可转债/宏观”
                    # 回退汇总阶段一科目二级分类。
                    secondary_label = ""
                    if actual is None:
                        secondary_actual, secondary_label = (
                            secondary_actual_for_file_strategy(
                                product.name,
                                plan.raw_strategy,
                                secondary_lookup,
                            )
                        )
                        if secondary_actual is not None:
                            actual = secondary_actual
                            resolved_category = secondary_label
                            stats.secondary_fallback_checks += 1
                            LOGGER.success(
                                "二级策略回退",
                                product.name,
                                secondary_label,
                                ppt_file.path.name,
                                f"一级分类未找到，按科目二级分类汇总实际占比{actual:.4f}",
                            )

                    if plan.extraction_status == "策略比例异常":
                        result = "策略比例异常"
                        stats.abnormal_ratio += 1
                        LOGGER.error(
                            "比例核查", product.name, plan.raw_strategy,
                            ppt_file.path.name,
                            f"第{plan.slide_number}页比例超出0%-100%或存在冲突，已停止比较",
                        )
                    elif not plan.primary_category and actual is None:
                        result = "文件策略无法识别，需人工确认"
                        stats.manual_review += 1
                        LOGGER.warning(
                            "比例核查", product.name, plan.raw_strategy,
                            ppt_file.path.name, result,
                        )
                    else:
                        result = compare_ratio(
                            actual,
                            plan.lower_limit,
                            plan.upper_limit,
                        )
                        if result == "符合":
                            stats.completed_checks += 1
                            stats.compliant += 1
                            LOGGER.success(
                                "比例核查", product.name, resolved_category,
                                ppt_file.path.name,
                                f"实际{actual:.4f}，计划{plan.lower_limit:.4f}-{plan.upper_limit:.4f}，符合",
                            )
                        elif result == "低于计划下限":
                            stats.completed_checks += 1
                            stats.below += 1
                            LOGGER.error(
                                "比例核查", product.name, resolved_category,
                                ppt_file.path.name,
                                f"实际{actual:.4f}，低于计划下限{plan.lower_limit:.4f}",
                            )
                        elif result == "高于计划上限":
                            stats.completed_checks += 1
                            stats.above += 1
                            LOGGER.error(
                                "比例核查", product.name, resolved_category,
                                ppt_file.path.name,
                                f"实际{actual:.4f}，高于计划上限{plan.upper_limit:.4f}",
                            )
                        elif result == "未找到对应产品策略":
                            stats.unmatched_actual_strategy += 1
                            LOGGER.info(
                                "比例核查", product.name, plan.raw_strategy,
                                ppt_file.path.name, result,
                            )
                        else:
                            stats.manual_review += 1
                            LOGGER.warning(
                                "比例核查", product.name, plan.raw_strategy,
                                ppt_file.path.name, result,
                            )

                    rows.append(CheckRow(
                        product_name=product.name,
                        primary_category=resolved_category,
                        file_name=ppt_file.path.name,
                        file_type=file_type,
                        file_strategy=plan.raw_strategy,
                        actual_total_ratio=actual,
                        lower_limit=plan.lower_limit,
                        upper_limit=plan.upper_limit,
                        check_result=result,
                        target_customer=metadata.target_customer,
                        product_positioning=metadata.product_positioning,
                        product_order=product.order,
                        source_order=plan.source_order,
                    ))

    rows.sort(
        key=lambda row: (
            row.product_order,
            type_order.get(row.file_type, 9),
            row.file_name.lower(),
            CATEGORY_ORDER.index(row.primary_category)
            if row.primary_category in CATEGORY_ORDER else len(CATEGORY_ORDER),
            row.source_order,
            row.file_strategy,
        )
    )
    return rows


# ============================================================
# 11. Excel输出
# ============================================================


OUTPUT_COLUMNS = (
    "产品名称",
    "一级分类",
    "文件名称",
    "文件类型",
    "文件策略",
    "产品策略总净值占比",
    "策略净值占比下限",
    "策略净值占比上限",
    "核查结果",
    "目标客户",
    "产品定位",
)


def rows_to_dataframe(rows: Sequence[CheckRow]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "产品名称": row.product_name,
                "一级分类": row.primary_category,
                "文件名称": row.file_name,
                "文件类型": row.file_type,
                "文件策略": row.file_strategy,
                "产品策略总净值占比": row.actual_total_ratio,
                "策略净值占比下限": row.lower_limit,
                "策略净值占比上限": row.upper_limit,
                "核查结果": row.check_result,
                "目标客户": row.target_customer,
                "产品定位": row.product_positioning,
            }
            for row in rows
        ],
        columns=OUTPUT_COLUMNS,
    )


def append_dataframe(worksheet: Any, dataframe: pd.DataFrame) -> None:
    worksheet.append(list(dataframe.columns))
    for row in dataframe.itertuples(index=False, name=None):
        worksheet.append(list(row))


def style_output(worksheet: Any) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    header_index = {clean_text(cell.value): cell.column for cell in worksheet[1]}
    for header in (
        "产品策略总净值占比",
        "策略净值占比下限",
        "策略净值占比上限",
    ):
        column = header_index.get(header)
        if column is None:
            continue
        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(row=row, column=column).number_format = "0.0000"

    result_column = header_index.get("核查结果")
    if result_column:
        red_results = {
            "低于计划下限",
            "高于计划上限",
            "策略比例异常",
            "未找到相关文件",
            "未找到策略配置比例",
            "未找到资产配置方案/投资方案页面",
            "PPT读取失败",
        }
        yellow_results = {
            "计划比例缺失",
            "文件策略无法识别，需人工确认",
            "配置方案未匹配，需人工确认",
            "配置方案存在歧义，需人工确认",
        }
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=result_column)
            result = clean_text(cell.value)
            if result in red_results:
                cell.fill = RED_FILL
            elif result in yellow_results:
                cell.fill = YELLOW_FILL
            # “未找到对应产品策略”明确不高亮。

    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for column in range(1, worksheet.max_column + 1):
        values = [
            clean_text(worksheet.cell(row=row, column=column).value)
            for row in range(1, min(worksheet.max_row, 500) + 1)
        ]
        max_length = max((len(value) for value in values), default=8)
        worksheet.column_dimensions[get_column_letter(column)].width = min(
            max(max_length + 2, 10), 42
        )


def save_workbook(dataframe: pd.DataFrame, output_path: Path, sheet_name: str) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    append_dataframe(worksheet, dataframe)
    style_output(worksheet)

    temporary_path = output_path.with_name(f"~${output_path.name}.tmp.xlsx")
    try:
        workbook.save(temporary_path)
        temporary_path.replace(output_path)
    finally:
        if temporary_path.exists():
            try:
                temporary_path.unlink()
            except OSError:
                pass


# ============================================================
# 12. 主流程
# ============================================================


def print_startup(config: Config) -> None:
    print("================ V5.2.1 阶段二：PPT文件提取与比例核查 ================")
    print(f"项目目录：{config.project_dir}")
    print(f"产品清单：{config.product_list_path}")
    print(f"阶段一结果：{config.stage1_output_path}")
    print(f"投决会目录：{config.committee_dir}")
    print(f"宣推材料目录：{config.promotion_dir}")
    print(f"输出文件：{config.output_path}")
    print("日期核验：关闭")
    print("===================================================================")


def print_summary(stats: RunStats, config: Config) -> None:
    print("\n================ V5.2.1 阶段二运行完成 ================")
    print(f"产品清单数量：{stats.product_count}")
    print(f"阶段一产品一级策略记录：{stats.actual_strategy_count}")
    print(f"阶段一科目二级分类记录：{stats.secondary_strategy_count}")
    print(f"投决会PPT数量：{stats.committee_file_count}")
    print(f"宣推PPT数量：{stats.promotion_file_count}")
    print(f"成功匹配PPT文件数量：{len(stats.matched_files)}")
    print(f"未匹配产品清单PPT数量：{stats.unmatched_files}")
    print(f"未找到投决会产品数量：{stats.products_without_committee}")
    print(f"未找到宣推文件产品数量：{stats.products_without_promotion}")
    print(f"提取策略配置数量：{stats.extracted_plans}")
    print(f"完成比例核查数量：{stats.completed_checks}")
    print(f"符合计划范围：{stats.compliant}")
    print(f"低于计划下限：{stats.below}")
    print(f"高于计划上限：{stats.above}")
    print(f"未找到对应产品策略：{stats.unmatched_actual_strategy}")
    print(f"二级分类回退核查：{stats.secondary_fallback_checks}")
    print(f"文本框补充提取策略：{stats.text_fallback_plans}")
    print(f"策略比例异常：{stats.abnormal_ratio}")
    print(f"需要人工确认：{stats.manual_review}")
    print(f"输出文件：{config.output_name}")
    print("======================================================")


def run_pipeline(config: Config = CONFIG) -> RunStats:
    print_startup(config)
    stats = RunStats()

    products = read_product_list(config.product_list_path, config.product_list_header)
    actual_records = read_actual_strategies(
        config.stage1_output_path,
        config.stage1_sheet,
    )
    secondary_records = read_secondary_strategies(
        config.stage1_output_path,
        config.stage1_detail_sheet,
    )
    stats.product_count = len(products)
    stats.actual_strategy_count = len(actual_records)
    stats.secondary_strategy_count = len(secondary_records)

    committee_files = scan_ppt_directory(config.committee_dir, "投决会文件")
    promotion_files = scan_ppt_directory(config.promotion_dir, "宣推文件")
    stats.committee_file_count = len(committee_files)
    stats.promotion_file_count = len(promotion_files)

    reader = PresentationReader()
    assignments = assign_ppt_files(
        [*committee_files, *promotion_files],
        products,
        reader,
        config,
        stats,
    )
    actual_lookup = build_actual_lookup(actual_records, products)
    secondary_lookup = build_secondary_lookup(
        secondary_records,
        products,
    )
    rows = build_check_rows(
        products,
        assignments,
        actual_lookup,
        secondary_lookup,
        reader,
        stats,
    )
    dataframe = rows_to_dataframe(rows)

    try:
        save_workbook(dataframe, config.output_path, config.output_sheet)
    except PermissionError as exc:
        LOGGER.error(
            "文件输出", "N/A", "N/A", config.output_name,
            "保存失败：文件可能正在被Excel占用",
        )
        raise PermissionError("输出文件可能正在被Excel占用") from exc
    except Exception as exc:
        LOGGER.error("文件输出", "N/A", "N/A", config.output_name, f"保存失败：{exc}")
        raise

    LOGGER.success("文件输出", "N/A", "N/A", config.output_name, "保存成功")
    print_summary(stats, config)
    return stats


def main() -> int:
    try:
        run_pipeline()
        return 0
    except KeyboardInterrupt:
        print(Fore.YELLOW + "程序已由用户中断。")
        return 130
    except Exception:
        print(Fore.RED + "程序运行失败，完整错误如下：")
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    sys.exit(main())
